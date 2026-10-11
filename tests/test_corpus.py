"""Independent publisher-reference validation, including explicit known failures."""

import importlib.util
import io
import json
from pathlib import Path

import openpyxl
import pytest

from tableproof.export import export_xlsx

ROOT = Path(__file__).resolve().parents[1]
CORPUS = ROOT / "samples" / "papers"
spec = importlib.util.spec_from_file_location("validate_corpus", ROOT / "scripts" / "validate_corpus.py")
validation_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(validation_module)

EXPECTED_SHAPES = {
    **{f"womack-2015-table-{n}": [11, 5] for n in range(1, 5)},
    **{f"womack-2015-table-{n}": [12, 4] for n in range(5, 9)},
    "womack-2015-table-9": [6, 6], "womack-2015-table-10": [6, 5],
    "womack-2015-table-11": [17, 6], "giofre-2017-table-2": [12, 7],
}


@pytest.fixture(scope="module")
def corpus():
    manifest = json.loads((CORPUS / "manifest.json").read_text(encoding="utf-8"))
    references = json.loads((CORPUS / "reference.json").read_text(encoding="utf-8"))
    return manifest, references, validation_module.validate(manifest, references)


def test_independent_reference_is_complete_and_attributed(corpus):
    manifest, references, _ = corpus
    assert len(manifest["tables"]) == len(references["tables"]) == 12
    for paper in references["papers"]:
        assert paper["xmlUrl"].startswith("https://journals.plos.org/")
        assert paper["xmlUrl"].endswith("type=manuscript")
        assert len(paper["xmlSha256"]) == 64
        assert paper["license"] == "CC BY 4.0"
    assert "no PDF extractor outputs" in references["referenceMethod"]
    for table in references["tables"]:
        assert [table["rowCount"], table["colCount"]] == EXPECTED_SHAPES[table["id"]]
        assert table["headerRowCount"] == (2 if table["id"] == "giofre-2017-table-2" else 1)
        assert table["keyValues"]
        assert table["caption"]
        assert table["cells"]
    womack11 = next(t for t in references["tables"] if t["id"] == "womack-2015-table-11")
    assert any("mathML" in c and not c["textComparisonEligible"] for c in womack11["cells"])


@pytest.mark.parametrize("table_id", EXPECTED_SHAPES)
def test_repaired_region_matches_distinct_official_key_values(corpus, table_id):
    _, _, findings = corpus
    result = next(t for t in findings["results"] if t["id"] == table_id)
    assert result["regionIdentityVerified"]
    assert all(k["presentInRegion"] for k in result["selectedRegionKeyChecks"])
    assert result["visualRegionReview"]["status"] == "reviewed"


@pytest.mark.parametrize("table_id", EXPECTED_SHAPES)
@pytest.mark.xfail(reason="Known automatic limitation: partial ruled grid and missing glyph mappings; manual correction is separate.", strict=False)
def test_automatic_structure_matches_independent_reference(corpus, table_id):
    _, _, findings = corpus
    result = next(t for t in findings["results"] if t["id"] == table_id)
    assert result["automaticError"] is None
    assert result["automatic"]["shapeMatch"]
    assert result["automatic"]["spansMatch"]
    assert result["automatic"]["exactCellTextMatches"] == result["automatic"]["eligibleReferenceCellCount"]


def test_unmapped_real_formula_pages_are_explicitly_blocked(corpus):
    _, _, findings = corpus
    for number in [9, 11]:
        result = next(t for t in findings["results"] if t["id"] == f"womack-2015-table-{number}")
        assert not result["mappingReliable"]
        assert result["automaticError"].startswith("UNRELIABLE_MAPPING")
        assert result["regionIdentityVerified"]


def test_partial_ruled_headers_are_retained_as_omitted_provenance(corpus):
    _, _, findings = corpus
    result = next(t for t in findings["results"] if t["id"] == "womack-2015-table-1")
    assert result["automatic"]["actualShape"] == [10, 3]
    assert not result["automatic"]["shapeMatch"]
    assert "TEXT_OUTSIDE_DETECTED_GRID" in result["automatic"]["warnings"]
    assert "rank" in result["automatic"]["omittedText"]
    assert "Journal" in result["automatic"]["omittedText"]


def test_real_complex_header_manual_workflow_matches_jats_and_exports(corpus):
    manifest, references, findings = corpus
    result = next(t for t in findings["results"] if t["id"] == "giofre-2017-table-2")
    assert result["manual"]["shapeMatch"]
    assert result["manual"]["spansMatch"]
    assert result["manual"]["exactCellTextMatches"] == result["manual"]["eligibleReferenceCellCount"] == 80
    assert result["manual"]["omittedText"] is None
    paper = next(p for p in manifest["papers"] if p["id"] == "giofre-2017")
    region = next(t for t in manifest["tables"] if t["id"] == "giofre-2017-table-2")
    path = CORPUS / paper["file"]
    pages = validation_module.parse_pdf(path, paper["sha256"])
    table = validation_module.manual_giofre(path, pages, paper["sha256"], region)
    project = {"id": "corpus-manual", "name": "Giofre manual review", "tables": [table],
               "fileHash": paper["sha256"], "filename": paper["file"], "version": 1,
               "history": [], "createdAt": "2026-10-11T00:00:00Z", "updatedAt": "2026-10-11T00:00:00Z"}
    workbook = openpyxl.load_workbook(io.BytesIO(export_xlsx(project)))
    sheet = workbook.worksheets[0]
    assert {str(r) for r in sheet.merged_cells.ranges} == {"B1:D1", "E1:G1"}
    reference = next(t for t in references["tables"] if t["id"] == table["id"] or t["id"] == "giofre-2017-table-2")
    for cell in reference["cells"]:
        actual = sheet.cell(cell["row"] + 1, cell["col"] + 1)
        assert validation_module.compact(str(actual.value or "")) == validation_module.compact(cell["text"])
        assert actual.data_type == "s"
    assert sheet.cell(12, 7).value == "0.04"
    assert workbook["Sources"].max_row > 80

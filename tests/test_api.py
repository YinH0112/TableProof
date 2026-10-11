import io
import json
import re
import threading
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from tableproof.api import create_app

ROOT = Path(__file__).resolve().parents[1]


def wait_job(client, job_id):
    for _ in range(200):
        result = client.get(f"/api/jobs/{job_id}").json()
        if result["status"] in {"complete", "failed", "cancelled"}:
            return result
        time.sleep(.02)
    raise AssertionError("Import timed out")


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path), base_url="http://127.0.0.1", raise_server_exceptions=False)


def import_demo(client):
    response = client.get("/api/demo")
    assert response.status_code == 202
    result = wait_job(client, response.json()["jobId"])
    assert result["status"] == "complete", result
    return client.get(f"/api/projects/{result['projectId']}").json()


def test_import_edit_persist_conflict_export_and_reextract(client):
    project = import_demo(client)
    project_id = project["id"]
    page = project["pages"][0]
    # Synthetic demo contains a ruled table; a full-page region selects it.
    response = client.post(f"/api/projects/{project_id}/tables", json={"version": project["version"], "page": 1, "bbox": page["mediaBox"]})
    assert response.status_code == 200, response.text
    project = response.json()
    table = project["tables"][0]
    cell = table["cells"][-1]
    refs = cell["sourceRefs"]
    original = cell["originalText"]
    old_version = project["version"]
    response = client.post(f"/api/projects/{project_id}/operations", json={"version": old_version, "operation": {"type": "edit", "tableId": table["id"], "cellId": cell["id"], "value": "=00123"}})
    assert response.status_code == 200, response.text
    project = response.json()
    edited = next(c for c in project["tables"][0]["cells"] if c["id"] == cell["id"])
    assert edited["value"] == "=00123" and edited["originalText"] == original
    assert edited["sourceRefs"] == refs and not edited["reviewed"]
    conflict = client.post(f"/api/projects/{project_id}/operations", json={"version": old_version, "operation": {"type": "rename", "name": "stale"}})
    assert conflict.status_code == 409
    reopened = TestClient(create_app(client.app.state.store.root), base_url="http://127.0.0.1")
    assert reopened.get(f"/api/projects/{project_id}").json() == project
    xlsx = client.get(f"/api/projects/{project_id}/export.xlsx")
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    anchor = workbook.worksheets[0].cell(edited["row"]+1, edited["col"]+1)
    assert anchor.value == "=00123" and anchor.data_type == "s"
    assert {"Sources", "Changes", "Metadata"} <= set(workbook.sheetnames)
    response = client.post(f"/api/projects/{project_id}/tables", json={"version": project["version"], "page": 1, "bbox": page["mediaBox"], "supersedes": table["id"]})
    assert response.status_code == 200
    project = response.json()
    assert len(project["tables"]) == 2
    assert project["tables"][0]["cells"][-1]["value"] == "=00123"
    assert project["tables"][1]["supersedes"] == table["id"]


def test_save_failure_does_not_acknowledge_edit(client, monkeypatch):
    project = import_demo(client)

    def fail(*_):
        raise OSError("No disk space")

    monkeypatch.setattr(client.app.state.store, "save", fail)
    response = client.post(f"/api/projects/{project['id']}/operations", json={"version": 1, "operation": {"type": "rename", "name": "unsaved"}})
    assert response.status_code == 500
    assert client.get(f"/api/projects/{project['id']}").json()["name"] == project["name"]


def test_invalid_and_encrypted_pdf_give_explicit_failure(client):
    response = client.post("/api/projects", files={"file": ("bad.pdf", b"not a pdf", "application/pdf")})
    failed = wait_job(client, response.json()["jobId"])
    assert failed["status"] == "failed" and failed["error"] == "INVALID_PDF"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    data = io.BytesIO()
    writer.write(data)
    response = client.post("/api/projects", files={"file": ("locked.pdf", data.getvalue(), "application/pdf")})
    failed = wait_job(client, response.json()["jobId"])
    assert failed["error"] == "ENCRYPTED_PDF"


def test_cancelled_import_does_not_create_project(client, monkeypatch):
    entered = threading.Event()

    def slow_parse(path, file_hash, progress_callback, cancelled):
        entered.set()
        for _ in range(200):
            if cancelled():
                raise ValueError("CANCELLED")
            time.sleep(.01)
        raise AssertionError("Not cancelled")

    monkeypatch.setattr("tableproof.api.parse_pdf", slow_parse)
    response = client.post("/api/projects", files={"file": ("slow.pdf", b"%PDF-fake", "application/pdf")})
    job_id = response.json()["jobId"]
    assert entered.wait(2)
    assert client.delete(f"/api/jobs/{job_id}").status_code == 200
    assert wait_job(client, job_id)["status"] == "cancelled"
    assert client.get("/api/projects").json() == []


def test_loopback_origin_and_regions(client):
    assert client.get("/api/health", headers={"host": "evil.example"}).status_code == 403
    assert client.get("/api/projects", headers={"origin": "https://evil.example"}).status_code == 403
    assert client.get("/api/projects/../../secrets").status_code == 404
    assert client.get("/api/health").status_code == 200
    project = import_demo(client)
    result = client.post(f"/api/projects/{project['id']}/tables", json={"version": 1, "page": 1, "bbox": [4, 4, 2, 2]})
    assert result.status_code == 422


def test_offline_pdf_module_worker_and_wasm_content_types(tmp_path):
    frontend = tmp_path / "frontend"
    frontend.mkdir()
    (frontend / "worker.mjs").write_text("export const test = true;", encoding="utf-8")
    (frontend / "parser.wasm").write_bytes(b"\0asm")
    browser = TestClient(create_app(tmp_path / "data", frontend), base_url="http://127.0.0.1")
    worker = browser.get("/worker.mjs")
    assert worker.status_code == 200
    assert worker.headers["content-type"].split(";")[0] == "text/javascript"
    assert worker.headers["x-content-type-options"] == "nosniff"
    assert browser.get("/parser.wasm").headers["content-type"] == "application/wasm"


def test_real_paper_manual_header_repair_edit_context_and_export(client):
    manifest = json.loads((ROOT / "samples/papers/manifest.json").read_text(encoding="utf-8"))
    references = json.loads((ROOT / "samples/papers/reference.json").read_text(encoding="utf-8"))
    region = next(t for t in manifest["tables"] if t["id"] == "giofre-2017-table-2")
    reference = next(t for t in references["tables"] if t["id"] == region["id"])
    data = (ROOT / "samples/papers/giofre-2017.pdf").read_bytes()
    response = client.post("/api/projects", files={"file": ("research.pdf", data, "application/pdf")})
    imported = wait_job(client, response.json()["jobId"])
    assert imported["status"] == "complete", imported
    project_id = imported["projectId"]
    project = client.get(f"/api/projects/{project_id}").json()
    response = client.post(f"/api/projects/{project_id}/tables", json={"version": project["version"], "page": region["page"], "bbox": region["bbox"], "columnLines": [142, 214, 287, 359, 432, 504]})
    assert response.status_code == 200, response.text
    project = response.json()
    table = project["tables"][0]
    for columns in [(1, 2, 3), (4, 5, 6)]:
        cell_ids = [c["id"] for c in table["cells"] if c["row"] == 0 and c["col"] in columns]
        response = client.post(f"/api/projects/{project_id}/operations", json={"version": project["version"], "operation": {"type": "merge", "tableId": table["id"], "cellIds": cell_ids}})
        assert response.status_code == 200, response.text
        project = response.json()
        table = project["tables"][0]
    assert [table["rowCount"], table["colCount"]] == [reference["rowCount"], reference["colCount"]]
    assert len(table["cells"]) == len(reference["cells"]) == 80
    expected = {(c["row"], c["col"]): c for c in reference["cells"]}
    for cell in table["cells"]:
        original = expected[(cell["row"], cell["col"])]
        assert re.sub(r"\s+", "", cell["value"]) == re.sub(r"\s+", "", original["text"])
        assert (cell["rowSpan"], cell["colSpan"]) == (original["rowSpan"], original["colSpan"])
    target = next(c for c in table["cells"] if c["row"] == 2 and c["col"] == 1)
    response = client.post(f"/api/projects/{project_id}/operations", json={"version": project["version"], "operation": {"type": "edit", "tableId": table["id"], "cellId": target["id"], "value": "00123"}})
    assert response.status_code == 200
    project = response.json()
    response = client.post(f"/api/projects/{project_id}/operations", json={"version": project["version"], "operation": {"type": "context", "tableId": table["id"], "kind": "caption", "page": 9, "bbox": [35, 696, 577, 730]}})
    assert response.status_code == 200, response.text
    project = response.json()
    assert project["tables"][0]["contexts"][0]["sourceRefs"][0]["charIds"]
    updated = next(c for c in project["tables"][0]["cells"] if c["id"] == target["id"])
    assert updated["originalText"] == target["originalText"] and updated["sourceRefs"] == target["sourceRefs"]
    workbook = load_workbook(io.BytesIO(client.get(f"/api/projects/{project_id}/export.xlsx").content))
    assert workbook.worksheets[0].cell(3, 2).value == "00123"
    for cell in project["tables"][0]["cells"]:
        assert (workbook.worksheets[0].cell(cell["row"]+1, cell["col"]+1).value or "") == cell["value"]
    assert client.get(f"/api/projects/{project_id}").json()["version"] == project["version"]

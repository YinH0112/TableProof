from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from tableproof.export import export_xlsx
from tableproof.operations import apply_operation, record_change, snapshot_project


def project():
    values = ["=SUM(1,2)", "0012", "123456789012345678901", "1E10", "+cmd|' /C calc'!A0", "−0.01", "μg\n95% CI"]
    cells = [{"id": f"c{i}", "row": 1, "col": i, "rowSpan": 1, "colSpan": 1,
              "originalText": value, "value": value, "reviewed": i == 0,
              "sourceRefs": [{"fileHash": "sha256", "page": 2, "bbox": [10, 20, 30, 40], "charIds": [f"p2-c{i}"]}]}
             for i, value in enumerate(values)]
    cells.insert(0, {"id": "header", "row": 0, "col": 0, "rowSpan": 1, "colSpan": 7,
                     "originalText": "Results", "value": "Results", "reviewed": False, "sourceRefs": []})
    context = {"id": "ctx", "kind": "unit", "text": "μg/mL", "sourceRefs": [{"fileHash": "sha256", "page": 2, "bbox": [11, 12, 13, 14], "charIds": ["p2-c9"]}]}
    return {"id": "p", "name": "Study", "fileHash": "sha256", "version": 3, "pages": [],
            "tables": [{"id": "t", "page": 2, "bbox": [0, 0, 100, 100], "rowCount": 2, "colCount": 7,
                        "cells": cells, "contexts": [context]}], "history": [], "undoStack": [], "redoStack": [],
            "createdAt": "2026-01-01T00:00:00Z", "updatedAt": "2026-01-02T00:00:00Z"}


def test_xlsx_preserves_literal_values_merged_header_context_and_source_evidence():
    initial = project()
    frozen = deepcopy(initial)
    data = export_xlsx(initial)
    assert initial == frozen
    workbook = load_workbook(BytesIO(data), data_only=False)
    assert workbook.sheetnames == ["Table 1", "Sources", "Changes", "Metadata"]
    sheet = workbook["Table 1"]
    assert "A1:G1" in sheet.merged_cells
    for index, item in enumerate(initial["tables"][0]["cells"][1:], 1):
        cell = sheet.cell(2, index)
        assert cell.value == item["value"]
        assert cell.data_type == "s" and cell.number_format == "@"
        assert "Sources" in cell.hyperlink.target
    assert sheet["B6"].value == "μg/mL"
    sources = list(workbook["Sources"].values)
    assert sources[2][8:11] == ("=SUM(1,2)", "=SUM(1,2)", "reviewed")
    assert sources[2][12:16] == ("sha256", "2", "[10,20,30,40]", '["p2-c0"]')
    assert sources[-1][1] == "unit" and sources[-1][9] == "μg/mL"
    metadata = dict(list(workbook["Metadata"].values)[1:])
    assert metadata["Pending cells"] == "7"
    # Inspect the serialized file itself, so safe load behavior cannot mask a
    # formula tag, precision conversion or leading-zero loss.
    with ZipFile(BytesIO(data)) as archive:
        for filename in archive.namelist():
            if filename.startswith("xl/worksheets/") and filename.endswith(".xml"):
                assert b"<f>" not in archive.read(filename)


def test_changes_include_undo_audit_and_original_source_after_reload():
    edited = apply_operation(project(), {"type": "edit", "tableId": "t", "cellId": "c1", "value": "0013"})
    undone = apply_operation(edited, {"type": "undo"})
    workbook = load_workbook(BytesIO(export_xlsx(undone)))
    rows = list(workbook["Changes"].values)
    assert rows[1][3] == "edit" and rows[2][3] == "undo"
    assert rows[1][7:9] == ("0012", "0013")
    assert rows[2][7:9] == ("0013", "0012")
    assert '"charIds":["p2-c1"]' in rows[1][11]
    assert workbook["Table 1"]["B2"].value == "0012"


def test_extraction_audit_expands_table_to_cells_and_contexts():
    initial = project()
    extracted = deepcopy(initial["tables"][0])
    initial["tables"] = []
    before = snapshot_project(initial)
    initial["tables"] = [extracted]
    record_change(initial, {"type": "extract", "tableId": "t"}, before)
    workbook = load_workbook(BytesIO(export_xlsx(initial)))
    rows = list(workbook["Changes"].values)[1:]
    assert sum(r[4] == "cell" for r in rows) == 8
    assert sum(r[4] == "context" for r in rows) == 1
    assert rows[1][8] == "=SUM(1,2)"
    assert rows[1][12].startswith('[{"fileHash":"sha256"')


@pytest.mark.parametrize("value,error", [("x" * 32768, "32,767"), ("value\x01", "control character")], ids=["oversize", "control"])
def test_export_refuses_to_silently_truncate_or_drop_unsupported_text(value, error):
    initial = project()
    initial["tables"][0]["cells"][1]["value"] = value
    with pytest.raises(ValueError, match=error):
        export_xlsx(initial)


def test_export_refuses_overlapping_cells():
    initial = project()
    initial["tables"][0]["cells"][1]["row"] = 0
    with pytest.raises(ValueError, match="Overlapping"):
        export_xlsx(initial)


def test_empty_project_still_exports_metadata_and_audit_sheets():
    initial = project()
    initial["tables"] = []
    workbook = load_workbook(BytesIO(export_xlsx(initial)))
    assert workbook.sheetnames == ["Sources", "Changes", "Metadata"]
    assert dict(list(workbook["Metadata"].values)[1:])["Tables"] == "0"


def test_superseded_table_versions_are_retained_and_explicitly_labelled():
    initial = project()
    replacement = {**deepcopy(initial["tables"][0]), "id": "t2", "supersedes": "t"}
    before = snapshot_project(initial)
    initial["tables"].append(replacement)
    record_change(initial, {"type": "extract", "tableId": "t2", "supersedes": "t"}, before)
    workbook = load_workbook(BytesIO(export_xlsx(initial)))
    assert workbook.sheetnames[:2] == ["Table 1 (superseded)", "Table 2"]
    assert workbook["Table 1 (superseded)"]["B2"].value == "0012"
    metadata = dict(list(workbook["Metadata"].values)[1:])
    assert metadata["Table 1 (superseded): status"] == "superseded"
    assert metadata["Table 2: status"] == "active"
    assert metadata["Table 2: supersedes"] == "t"
    assert metadata["Table 2: extraction save version"] == "4"


def test_unassigned_selected_text_is_visible_and_source_linked_without_changing_grid():
    initial = project()
    table = initial["tables"][0]
    text = '=HYPERLINK("https://example.invalid", "Table caption")\nUnits: μg/mL'
    refs = [{"fileHash": "sha256", "page": 2, "bbox": [5, 150, 500, 180],
             "charIds": ["p2-caption-0", "p2-caption-1"]}]
    table["omittedText"] = {"text": text, "sourceRefs": refs}
    table["extractionStrategy"] = "lines"
    table["warnings"] = ["TEXT_OUTSIDE_DETECTED_GRID"]
    table["mappingReliable"] = False
    frozen = deepcopy(initial)
    data = export_xlsx(initial)
    assert initial == frozen
    workbook = load_workbook(BytesIO(data), data_only=False)
    sheet = workbook["Table 1"]
    assert sheet["A8"].value == "Selected text outside grid"
    assert sheet["A9"].value == text and sheet["A9"].data_type == "s"
    assert "A9:G9" in sheet.merged_cells
    sources = list(workbook["Sources"].values)
    omitted = sources[-1]
    assert omitted[1] == "omittedText" and omitted[2] == "t-omitted-text"
    assert omitted[3] == "'Table 1'!A9"
    assert omitted[8:11] == (text, text, "unassigned")
    assert omitted[12:16] == ("sha256", "2", "[5,150,500,180]", '["p2-caption-0","p2-caption-1"]')
    assert sheet["A9"].hyperlink.target == f"#'Sources'!A{len(sources)}"
    assert sheet["A1"].value == "Results"
    for item in table["cells"][1:]:
        assert sheet.cell(item["row"] + 1, item["col"] + 1).value == item["value"]
    metadata = dict(list(workbook["Metadata"].values)[1:])
    assert metadata["Table 1: extraction strategy"] == "lines"
    assert metadata["Table 1: warnings"] == '["TEXT_OUTSIDE_DETECTED_GRID"]'
    assert metadata["Table 1: mapping reliable"] == "false"
    assert "not a corrected grid value" in metadata["Selected text outside grid"]
    with ZipFile(BytesIO(data)) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")


def test_unassigned_text_stays_in_extraction_audit_and_undo_export():
    initial = project()
    extracted = deepcopy(initial["tables"][0])
    extracted["omittedText"] = {"text": "Unassigned header", "sourceRefs": [{
        "fileHash": "sha256", "page": 2, "bbox": [5, 150, 500, 180], "charIds": ["p2-caption-0"]}]}
    initial["tables"] = []
    before = snapshot_project(initial)
    initial["tables"] = [extracted]
    record_change(initial, {"type": "extract", "tableId": "t"}, before)
    undone = apply_operation(initial, {"type": "undo"})
    workbook = load_workbook(BytesIO(export_xlsx(undone)))
    changes = list(workbook["Changes"].values)[1:]
    table_rows = [r for r in changes if r[4] == "table"]
    assert '"text":"Unassigned header"' in table_rows[0][14]
    assert '"charIds":["p2-caption-0"]' in table_rows[0][14]
    assert '"text":"Unassigned header"' in table_rows[1][13]

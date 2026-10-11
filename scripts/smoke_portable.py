"""Verify the portable server with Python/Node removed from its PATH."""
import io
import json
import os
import subprocess
import tempfile
import time
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
URL = "http://127.0.0.1:18766"


def request(path, body=None):
    req = Request(URL + path, data=json.dumps(body).encode() if body else None,
                  headers={"Content-Type": "application/json"} if body else {})
    with urlopen(req, timeout=15) as response:
        return json.loads(response.read())


def main():
    env = dict(os.environ)
    env["PATH"] = str(Path(env.get("SystemRoot", "C:/Windows")) / "System32")
    env.pop("PYTHONPATH", None)
    env.pop("PYTHONHOME", None)
    with tempfile.TemporaryDirectory(prefix="tableproof-portable-") as temp:
        log_path = Path(temp) / "server.log"
        with log_path.open("wb") as log:
            server = subprocess.Popen([str(ROOT / "dist" / "TableProof" / "TableProof.exe"), "--no-browser", "--port", "18766", "--data-dir", str(Path(temp) / "data")], cwd=temp, env=env, stdout=log, stderr=log)
            try:
                for _ in range(150):
                    if server.poll() is not None:
                        raise RuntimeError(log_path.read_text(errors="replace"))
                    try:
                        request("/api/health")
                        break
                    except OSError:
                        time.sleep(.1)
                else:
                    raise RuntimeError("Portable server did not become ready")
                with urlopen(URL, timeout=10) as response:
                    assert b'<html' in response.read()
                bundled_assets = ROOT / "dist" / "TableProof" / "_internal" / "frontend" / "dist" / "assets"
                worker = next(bundled_assets.glob("pdf.worker*.mjs"))
                with urlopen(URL + "/assets/" + worker.name, timeout=10) as response:
                    assert response.headers.get_content_type() == "text/javascript"
                    assert response.read(100)
                job = request("/api/demo")
                for _ in range(200):
                    result = request(f"/api/jobs/{job['jobId']}")
                    if result["status"] not in {"pending", "running"}:
                        break
                    time.sleep(.05)
                assert result["status"] == "complete", result
                project_id = result["projectId"]
                project = request(f"/api/projects/{project_id}")
                project = request(f"/api/projects/{project_id}/tables", {"version": project["version"], "page": 1, "bbox": [50, 510, 550, 710]})
                table, cell = project["tables"][0], project["tables"][0]["cells"][-1]
                project = request(f"/api/projects/{project_id}/operations", {"version": project["version"], "operation": {"type": "edit", "tableId": table["id"], "cellId": cell["id"], "value": "=00123"}})
                assert request(f"/api/projects/{project_id}")["version"] == project["version"]
                with urlopen(URL + f"/api/projects/{project_id}/export.xlsx", timeout=10) as response:
                    workbook = load_workbook(io.BytesIO(response.read()))
                anchor = workbook.worksheets[0].cell(cell["row"]+1, cell["col"]+1)
                assert anchor.value == "=00123" and anchor.data_type == "s"
                assert {"Sources", "Changes", "Metadata"} <= set(workbook.sheetnames)
                print("Portable passed: isolated PATH, HTML, sample parse, extract, edit, persist, string-safe XLSX.")
            finally:
                server.terminate()
                try:
                    server.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait()


if __name__ == "__main__":
    main()

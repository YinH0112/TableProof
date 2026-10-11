"""Export reviewed tables together with source evidence and the audit trail."""

from __future__ import annotations

from io import BytesIO
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="173C35")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def _write(sheet, row: int, col: int, value: Any):
    text = _text(value)
    if len(text) > 32767:
        raise ValueError("Excel cannot store a field longer than 32,767 characters; no text was truncated.")
    if ILLEGAL_CHARACTERS_RE.search(text) or any(0xD800 <= ord(c) <= 0xDFFF or ord(c) in (0xFFFE, 0xFFFF) for c in text):
        raise ValueError("Excel cannot store a control character in this text; export was cancelled.")
    cell = sheet.cell(row=row, column=col)
    cell.value = text
    # Assignment auto-detects '=…' as a formula. Override it deliberately rather
    # than prefixing an apostrophe, which would alter the researcher’s value.
    cell.data_type = "s"
    cell.number_format = "@"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _append(sheet, values: list[Any]) -> int:
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for col, value in enumerate(values, 1):
        _write(sheet, row, col, value)
    return row


def _headers(sheet, columns: list[str]) -> None:
    _append(sheet, columns)
    for cell in sheet[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    sheet.freeze_panes = "A2"


def _ref_rows(sheet, table: dict, item: dict, kind: str, address: str) -> int:
    first = sheet.max_row + 1
    refs = item.get("sourceRefs", []) or [{}]
    for index, ref in enumerate(refs, 1):
        _append(sheet, [table["id"], kind, item["id"], address,
                        item.get("row", ""), item.get("col", ""),
                        item.get("rowSpan", ""), item.get("colSpan", ""),
                        item.get("originalText", ""), item.get("value", item.get("text", "")),
                        "reviewed" if item.get("reviewed") else "pending" if kind == "cell"
                        else "unassigned" if kind == "omittedText" else "context",
                        index if ref else "", ref.get("fileHash", ""), ref.get("page", ""),
                        ref.get("bbox", ""), ref.get("charIds", ""),
                        "" if ref else "No captured source characters"])
    return first


def _change_rows(history: list[dict]):
    for record in history:
        action = record.get("action", {"type": record.get("type", "")})
        changes = record.get("changes", [])
        if not changes:
            yield record, action, {}
        for change in changes:
            # Extraction records store a whole newly added table. Expand the
            # cells and contexts instead of putting a large JSON blob in a field.
            if change.get("entity") == "table":
                before, after = change.get("before") or {}, change.get("after") or {}
                if isinstance(before, dict) and isinstance(after, dict) and ("cells" in before or "cells" in after):
                    old_cells = {c["id"]: c for c in before.get("cells", [])}
                    new_cells = {c["id"]: c for c in after.get("cells", [])}
                    for cell_id in dict.fromkeys([*old_cells, *new_cells]):
                        yield record, action, {"entity": "cell", "tableId": change.get("tableId"),
                                               "cellId": cell_id, "before": old_cells.get(cell_id),
                                               "after": new_cells.get(cell_id)}
                    old_contexts = {c["id"]: c for c in before.get("contexts", [])}
                    new_contexts = {c["id"]: c for c in after.get("contexts", [])}
                    for context_id in dict.fromkeys([*old_contexts, *new_contexts]):
                        yield record, action, {"entity": "context", "tableId": change.get("tableId"),
                                               "contextId": context_id, "before": old_contexts.get(context_id),
                                               "after": new_contexts.get(context_id)}
                    change = {**change, "before": {k: v for k, v in before.items() if k not in ("cells", "contexts")},
                              "after": {k: v for k, v in after.items() if k not in ("cells", "contexts")}}
            yield record, action, change


def _change_value(value: Any) -> Any:
    if isinstance(value, dict):
        return value.get("value", value.get("text", ""))
    return value


def export_xlsx(project: dict[str, Any]) -> bytes:
    """Create XLSX bytes. All exported fields are literal strings, never formulas.

    The Sources sheet uses zero-based grid indices and original PDF user-space
    bounding boxes. Pending review is retained and disclosed in the workbook.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    sources = workbook.create_sheet("Sources")
    _headers(sources, ["Table ID", "Kind", "Item ID", "Workbook address", "Row (0-based)", "Col (0-based)",
                       "Row span", "Col span", "Original text", "Current value / context", "Review status",
                       "Source reference", "PDF SHA-256", "Page (1-based)", "BBox (PDF user space)",
                       "Character IDs", "Evidence note"])
    tables = project.get("tables", [])
    superseded_ids = {t.get("supersedes") for t in tables if t.get("supersedes")}
    table_metadata = []
    for table_index, table in enumerate(tables, 1):
        status = "superseded" if table["id"] in superseded_ids else "active"
        title = f"Table {table_index}" + (" (superseded)" if status == "superseded" else "")
        sheet = workbook.create_sheet(title, table_index - 1)
        extraction = next((h for h in project.get("history", [])
                           if h.get("type") == "extract" and h.get("action", {}).get("tableId") == table["id"]), {})
        table_metadata.extend([(f"{title}: ID", table["id"]), (f"{title}: status", status),
                               (f"{title}: page", table.get("page", "")),
                               (f"{title}: supersedes", table.get("supersedes", "")),
                               (f"{title}: extraction save version", extraction.get("version", "")),
                               (f"{title}: extraction strategy", table.get("extractionStrategy", "")),
                               (f"{title}: warnings", table.get("warnings", [])),
                               (f"{title}: mapping reliable", table.get("mappingReliable", ""))])
        row_count, col_count = table["rowCount"], table["colCount"]
        if row_count < 1 or col_count < 1 or row_count > 1048576 or col_count > 16384:
            raise ValueError("Table dimensions exceed the Excel worksheet limits.")
        contexts = table.get("contexts", [])
        context_end = row_count + 3 + len(contexts)
        omitted = table.get("omittedText")
        if contexts and context_end > 1048576:
            raise ValueError("Table context exceeds the Excel worksheet row limit.")
        omitted_row = context_end + 2 if contexts else row_count + 3
        if omitted and omitted_row + 1 > 1048576:
            raise ValueError("Selected text outside the grid exceeds the Excel worksheet row limit.")
        sheet.freeze_panes = "A2"
        occupied = set()
        for item in sorted(table.get("cells", []), key=lambda c: (c["row"], c["col"])):
            r0, c0 = item["row"], item["col"]
            rows, cols = item.get("rowSpan", 1), item.get("colSpan", 1)
            if min(r0, c0) < 0 or min(rows, cols) < 1 or r0 + rows > row_count or c0 + cols > col_count:
                raise ValueError("A cell is outside the table grid.")
            slots = {(r, c) for r in range(r0, r0 + rows) for c in range(c0, c0 + cols)}
            if occupied & slots:
                raise ValueError("Overlapping cells cannot be exported without losing a value.")
            occupied |= slots
            address = f"'{sheet.title}'!{get_column_letter(c0 + 1)}{r0 + 1}"
            source_row = _ref_rows(sources, table, item, "cell", address)
            cell = _write(sheet, r0 + 1, c0 + 1, item.get("value", ""))
            cell.hyperlink = f"#'Sources'!A{source_row}"
            cell.comment = Comment(f"Review: {'reviewed' if item.get('reviewed') else 'pending'}\n"
                                   f"Original text: {item.get('originalText', '')}\n"
                                   f"Sources: row {source_row} on the Sources sheet.", "TableProof")
            if not item.get("reviewed"):
                cell.fill = PatternFill("solid", fgColor="FFF0CF")
            if rows > 1 or cols > 1:
                sheet.merge_cells(start_row=r0 + 1, start_column=c0 + 1,
                                  end_row=r0 + rows, end_column=c0 + cols)
        context_row = row_count + 3
        if contexts:
            _write(sheet, context_row, 1, "Table context")
            for item in contexts:
                context_row += 1
                address = f"'{sheet.title}'!B{context_row}"
                _write(sheet, context_row, 1, item.get("kind", ""))
                context_cell = _write(sheet, context_row, 2, item.get("text", ""))
                source_row = _ref_rows(sources, table, item, item.get("kind", "context"), address)
                context_cell.hyperlink = f"#'Sources'!A{source_row}"
        if omitted:
            # These original characters were inside the user's selection but
            # outside the detected cells. Retain them as evidence, never as a
            # corrected grid value or an inferred header.
            heading = _write(sheet, omitted_row, 1, "Selected text outside grid")
            heading.font = Font(bold=True, color="865B21")
            heading.fill = PatternFill("solid", fgColor="FFF0CF")
            item = {**omitted, "id": f"{table['id']}-omitted-text", "originalText": omitted.get("text", "")}
            address = f"'{sheet.title}'!A{omitted_row + 1}"
            source_row = _ref_rows(sources, table, item, "omittedText", address)
            retained = _write(sheet, omitted_row + 1, 1, omitted.get("text", ""))
            retained.hyperlink = f"#'Sources'!A{source_row}"
            retained.fill = PatternFill("solid", fgColor="FFF8E8")
            retained.comment = Comment("Original selected text not assigned to a grid cell. "
                                       "Inspect its Sources reference and review the selection or grid.", "TableProof")
            if col_count > 1:
                sheet.merge_cells(start_row=omitted_row + 1, start_column=1,
                                  end_row=omitted_row + 1, end_column=col_count)
            # Merged cells do not reliably auto-size their height in Excel.
            lines = sum(max(1, (len(line) + col_count * 20 - 1) // (col_count * 20))
                        for line in omitted.get("text", "").split("\n"))
            sheet.row_dimensions[omitted_row + 1].height = min(300, max(30, lines * 15))
        for col in range(1, col_count + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22
    changes = workbook.create_sheet("Changes")
    _headers(changes, ["Audit ID", "Timestamp (UTC)", "Save version", "Action", "Entity", "Table ID", "Item ID",
                       "Before value", "After value", "Before reviewed", "After reviewed",
                       "Before evidence", "After evidence", "Before details", "After details", "Operation"])
    for record, action, change in _change_rows(project.get("history", [])):
        before, after = change.get("before"), change.get("after")
        _append(changes, [record.get("id", ""), record.get("createdAt", ""), record.get("version", ""),
                         record.get("type", action.get("type", "")), change.get("entity", ""),
                         change.get("tableId", ""), change.get("cellId", change.get("contextId", "")),
                         _change_value(before), _change_value(after),
                         before.get("reviewed", "") if isinstance(before, dict) else "",
                         after.get("reviewed", "") if isinstance(after, dict) else "",
                         before.get("sourceRefs", []) if isinstance(before, dict) else "",
                         after.get("sourceRefs", []) if isinstance(after, dict) else "", before, after, action])
    metadata = workbook.create_sheet("Metadata")
    _headers(metadata, ["Field", "Value"])
    cells = [c for table in project.get("tables", []) for c in table.get("cells", [])]
    fields = [("Tool", "TableProof Research"), ("Project ID", project.get("id", "")),
              ("Project name", project.get("name", "")), ("Project version", project.get("version", "")),
              ("PDF SHA-256", project.get("fileHash", "")), ("Created at", project.get("createdAt", "")),
              ("Saved at", project.get("updatedAt", "")), ("Tables", len(project.get("tables", []))),
              ("Cells", len(cells)), ("Reviewed cells", sum(bool(c.get("reviewed")) for c in cells)),
              ("Pending cells", sum(not c.get("reviewed") for c in cells)),
              ("Value encoding", "Literal text; formulas are not executed; leading zeros are retained"),
              ("Grid coordinates", "Sources row/col indices are zero-based; worksheet addresses are one-based"),
              ("Source coordinates", "Original PDF user space: [x0,y0,x1,y1]; page numbers are one-based"),
              ("Review status", "Pending cells are highlighted; exporting does not mark them reviewed")]
    fields.extend([("Selected text outside grid", "Original unassigned selection text is retained separately "
                   "below each table and in Sources with kind omittedText; it is not a corrected grid value")])
    fields.extend([("Active tables", sum(t["id"] not in superseded_ids for t in tables)),
                   ("Superseded tables", sum(t["id"] in superseded_ids for t in tables)),
                   ("Table versions", "All versions are retained; sheets labelled superseded have been replaced by a later selection")])
    for field in fields + table_metadata:
        _append(metadata, list(field))
    for sheet in (sources, changes, metadata):
        sheet.auto_filter.ref = sheet.dimensions
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 28 if col < 9 else 42
    workbook.active = 0
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
